import csv
import json
import math
import os
import re
import subprocess
import sys
import traceback
from bisect import bisect_left
from datetime import datetime, date, time, timedelta
from pathlib import Path

from PySide6.QtCore import Qt, QAbstractTableModel, QModelIndex
from PySide6.QtGui import QAction
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QPushButton,
    QLabel, QFileDialog, QTableView, QTextEdit, QMessageBox, QSplitter,
    QProgressBar, QCheckBox, QDoubleSpinBox, QSpinBox
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except Exception:
    Workbook = None

APP_TITLE = "PEF EXIF + NMEA Explorer"

BASE_COLUMNS = [
    "ファイル名", "フルパス", "撮影時刻", "緯度", "経度", "標高", "方位",
    "ピッチ", "ロール", "レンズ", "焦点距離", "シャッター", "絞り", "ISO",
    "カメラ機種", "露出補正", "ホワイトバランス", "撮影モード", "RAW全タグ数"
]

NMEA_COLUMNS = [
    "NMEA照合", "NMEA時刻", "時刻差秒",
    "NMEA緯度", "NMEA経度", "NMEA標高", "水平距離差m", "標高差m",
    "GGA測位品質", "GGA衛星数", "GGA_HDOP", "GGAジオイド高", "GGA補正経過秒", "GGA基準局ID",
    "PRDID_1", "PRDID_2", "PRDID_3", "PRDID文"
]

OUTPUT_COLUMNS = BASE_COLUMNS + NMEA_COLUMNS

TAG_CANDIDATES = {
    "撮影時刻": ["DateTimeOriginal", "CreateDate", "ModifyDate", "SubSecDateTimeOriginal"],
    "緯度": ["GPSLatitude"],
    "経度": ["GPSLongitude"],
    "標高": ["GPSAltitude", "Altitude"],
    "方位": ["GPSImgDirection", "GPSDestBearing", "ImageDirection", "CompassHeading", "Heading", "Azimuth"],
    "ピッチ": ["Pitch", "PitchAngle", "CameraPitch", "LevelPitch", "AccelerometerPitch"],
    "ロール": ["Roll", "RollAngle", "CameraRoll", "LevelRoll", "AccelerometerRoll"],
    "レンズ": ["LensModel", "LensType", "LensID", "Lens", "LensInfo"],
    "焦点距離": ["FocalLength", "FocalLengthIn35mmFormat"],
    "シャッター": ["ExposureTime", "ShutterSpeed", "ShutterSpeedValue"],
    "絞り": ["FNumber", "Aperture", "ApertureValue"],
    "ISO": ["ISO", "SensitivityType", "RecommendedExposureIndex"],
    "カメラ機種": ["Model", "CameraModelName", "PentaxModelID"],
    "露出補正": ["ExposureCompensation", "ExposureCompensationSet", "ExposureBiasValue"],
    "ホワイトバランス": ["WhiteBalance", "WhiteBalanceMode", "WBMode"],
    "撮影モード": ["ExposureProgram", "ExposureMode", "ShootingMode", "DriveMode", "PictureMode"]
}


def app_base_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent.parent


def find_exiftool() -> str:
    # 重要: PyInstallerの一時展開先(_MEIPASS)に同梱したExifToolは、
    # Perl標準モジュールの参照に失敗する場合があるため最初には使わない。
    # EXEまたはプロジェクトフォルダ横の tools\exiftool.exe を優先する。
    base = app_base_dir()
    candidates = [
        base / "tools" / "exiftool.exe",
        base / "exiftool.exe",
        Path.cwd() / "tools" / "exiftool.exe",
        Path.cwd() / "exiftool.exe",
    ]
    for c in candidates:
        if c.exists():
            return str(c)
    # 最後の保険として、PATHと_MEIPASSを確認
    for p in os.environ.get("PATH", "").split(os.pathsep):
        cand = Path(p) / "exiftool.exe"
        if cand.exists():
            return str(cand)
    base_tmp = getattr(sys, "_MEIPASS", None)
    if base_tmp:
        cand = Path(base_tmp) / "exiftool.exe"
        if cand.exists():
            return str(cand)
    raise FileNotFoundError("exiftool.exe が見つかりません。EXEと同じフォルダの tools\\exiftool.exe に配置してください。")


def run_exiftool(files, numeric=False):
    exiftool = find_exiftool()
    args = [exiftool, "-j", "-G1", "-a", "-u", "-charset", "filename=utf8"]
    if numeric:
        args.append("-n")
    args.extend([str(Path(f)) for f in files])
    startupinfo = None
    if os.name == "nt":
        startupinfo = subprocess.STARTUPINFO()
        startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
    cp = subprocess.run(
        args,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
        startupinfo=startupinfo,
        shell=False
    )
    if cp.returncode not in (0, 1):
        raise RuntimeError(cp.stderr.strip() or "ExifTool実行エラー")
    if not cp.stdout.strip():
        if cp.stderr.strip():
            raise RuntimeError(cp.stderr.strip())
        return []
    return json.loads(cp.stdout)


def normalize_key(k: str) -> str:
    return re.sub(r"[^a-z0-9]", "", k.lower().split(":")[-1])


def get_tag(data: dict, candidates):
    norm_map = {normalize_key(k): v for k, v in data.items()}
    for cand in candidates:
        key = normalize_key(cand)
        if key in norm_map:
            return norm_map[key]
    for cand in candidates:
        key = normalize_key(cand)
        for nk, value in norm_map.items():
            if nk.endswith(key) or key.endswith(nk):
                return value
    return ""


def to_number(v):
    if v is None or v == "":
        return ""
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if m:
        try:
            return float(m.group(0))
        except Exception:
            return v
    return v


def parse_exif_datetime(v):
    if not v:
        return None
    s = str(v).strip()
    s = s.split("+")[0].split("Z")[0]
    forms = ["%Y:%m:%d %H:%M:%S.%f", "%Y:%m:%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S"]
    for fmt in forms:
        try:
            return datetime.strptime(s[:26], fmt)
        except Exception:
            pass
    return None


def build_row(print_data: dict, num_data: dict) -> dict:
    src = print_data.get("SourceFile") or num_data.get("SourceFile") or ""
    row = {c: "" for c in OUTPUT_COLUMNS}
    row["ファイル名"] = Path(src).name
    row["フルパス"] = str(Path(src)) if src else ""
    for col in BASE_COLUMNS:
        if col in ("ファイル名", "フルパス", "RAW全タグ数"):
            continue
        if col in ("緯度", "経度", "標高", "方位", "ピッチ", "ロール"):
            v = get_tag(num_data, TAG_CANDIDATES.get(col, []))
            if v == "":
                v = get_tag(print_data, TAG_CANDIDATES.get(col, []))
            row[col] = to_number(v)
        else:
            row[col] = get_tag(print_data, TAG_CANDIDATES.get(col, []))
    row["RAW全タグ数"] = len(print_data)
    return row


def collect_pef_paths(paths):
    files = []
    for p in paths:
        path = Path(p)
        if path.is_file() and path.suffix.lower() == ".pef":
            files.append(path)
        elif path.is_dir():
            files.extend(sorted(path.rglob("*.PEF")))
            files.extend(sorted(path.rglob("*.pef")))
    seen = set()
    unique = []
    for f in files:
        key = os.path.normcase(str(f.absolute()))
        if key not in seen:
            unique.append(f)
            seen.add(key)
    return unique


def nmea_checksum_ok(line):
    line = line.strip()
    if not line.startswith("$") or "*" not in line:
        return True
    body, cs = line[1:].split("*", 1)
    cs = cs.strip()[:2]
    val = 0
    for ch in body:
        val ^= ord(ch)
    try:
        return val == int(cs, 16)
    except Exception:
        return False


def parse_latlon_nmea(value, hemi):
    if not value:
        return ""
    try:
        dot = value.find(".")
        deg_len = dot - 2
        deg = float(value[:deg_len])
        minutes = float(value[deg_len:])
        dec = deg + minutes / 60.0
        if hemi in ("S", "W"):
            dec *= -1
        return dec
    except Exception:
        return ""


def parse_utc_hhmmss(v):
    if not v:
        return None
    try:
        hh = int(v[0:2])
        mm = int(v[2:4])
        ss = float(v[4:])
        sec = int(ss)
        micro = int(round((ss - sec) * 1_000_000))
        return time(hh, mm, sec, micro)
    except Exception:
        return None


def guess_date_from_filename(path):
    name = Path(path).name
    m = re.search(r"(20\d{2})[-_](\d{2})[-_](\d{2})", name)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def parse_nmea_file(path, utc_offset_hours=9, fallback_date=None):
    base_date = guess_date_from_filename(path) or fallback_date or date.today()
    records = []
    bad_checksum = 0
    last_prdid = None
    current_idx = None
    prev_dt = None
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        for line_no, raw in enumerate(f, start=1):
            line = raw.strip()
            if not line:
                continue
            if not nmea_checksum_ok(line):
                bad_checksum += 1
            body = line[1:].split("*", 1)[0] if line.startswith("$") else line.split("*", 1)[0]
            parts = body.split(",")
            typ = parts[0]
            if typ.endswith("GGA") and len(parts) >= 15:
                t = parse_utc_hhmmss(parts[1])
                if t is None:
                    continue
                dt_utc = datetime.combine(base_date, t)
                if prev_dt is not None and dt_utc < prev_dt - timedelta(hours=12):
                    dt_utc += timedelta(days=1)
                prev_dt = dt_utc
                dt_local = dt_utc + timedelta(hours=utc_offset_hours)
                rec = {
                    "dt_local": dt_local,
                    "nmea_type": typ,
                    "lat": parse_latlon_nmea(parts[2], parts[3]),
                    "lon": parse_latlon_nmea(parts[4], parts[5]),
                    "quality": parts[6],
                    "satellites": parts[7],
                    "hdop": to_number(parts[8]),
                    "altitude": to_number(parts[9]),
                    "geoid_sep": to_number(parts[11]) if len(parts) > 11 else "",
                    "dgps_age": to_number(parts[13]) if len(parts) > 13 else "",
                    "station_id": parts[14] if len(parts) > 14 else "",
                    "prdid_1": "", "prdid_2": "", "prdid_3": "", "prdid_sentence": "",
                    "line": line_no,
                }
                if last_prdid:
                    rec.update(last_prdid)
                records.append(rec)
                current_idx = len(records) - 1
            elif typ == "PRDID" and len(parts) >= 3:
                vals = [p for p in parts[1:]]
                pr = {
                    "prdid_1": to_number(vals[0]) if len(vals) > 0 else "",
                    "prdid_2": to_number(vals[1]) if len(vals) > 1 else "",
                    "prdid_3": to_number(vals[2]) if len(vals) > 2 else "",
                    "prdid_sentence": line,
                }
                last_prdid = pr
                if current_idx is not None:
                    records[current_idx].update(pr)
    records.sort(key=lambda r: r["dt_local"])
    return records, bad_checksum


def haversine_m(lat1, lon1, lat2, lon2):
    try:
        lat1, lon1, lat2, lon2 = map(float, [lat1, lon1, lat2, lon2])
    except Exception:
        return ""
    r = 6371008.8
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp/2)**2 + math.cos(p1) * math.cos(p2) * math.sin(dl/2)**2
    return 2 * r * math.atan2(math.sqrt(a), math.sqrt(1-a))


def nearest_nmea_record(records, target_dt, threshold_sec):
    if not records or target_dt is None:
        return None, ""
    times = [r["dt_local"] for r in records]
    i = bisect_left(times, target_dt)
    candidates = []
    if i < len(records):
        candidates.append(records[i])
    if i > 0:
        candidates.append(records[i-1])
    if not candidates:
        return None, ""
    best = min(candidates, key=lambda r: abs((r["dt_local"] - target_dt).total_seconds()))
    diff = (best["dt_local"] - target_dt).total_seconds()
    if abs(diff) <= threshold_sec:
        return best, diff
    return None, diff


class DropArea(QLabel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAcceptDrops(True)
        self.setAlignment(Qt.AlignCenter)
        self.setText("ここに PEF ファイル / フォルダをドラッグ＆ドロップ\nまたは［ファイル追加］［フォルダ追加］を使用")
        self.setStyleSheet("QLabel { border: 2px dashed #777; padding: 24px; font-size: 14px; background: #fafafa; }")
        self.on_files = None

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self, event):
        paths = [u.toLocalFile() for u in event.mimeData().urls()]
        if self.on_files:
            self.on_files(paths)


class ResultModel(QAbstractTableModel):
    def __init__(self):
        super().__init__()
        self.rows = []

    def rowCount(self, parent=QModelIndex()):
        return len(self.rows)

    def columnCount(self, parent=QModelIndex()):
        return len(OUTPUT_COLUMNS)

    def data(self, index, role=Qt.DisplayRole):
        if not index.isValid() or role not in (Qt.DisplayRole, Qt.EditRole):
            return None
        return str(self.rows[index.row()].get(OUTPUT_COLUMNS[index.column()], ""))

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        if role != Qt.DisplayRole:
            return None
        if orientation == Qt.Horizontal:
            return OUTPUT_COLUMNS[section]
        return section + 1

    def set_rows(self, rows):
        self.beginResetModel()
        self.rows = rows
        self.endResetModel()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_TITLE)
        self.resize(1500, 850)
        self.model = ResultModel()
        self.raw_print = []
        self.raw_numeric = []
        self.rows = []
        self.nmea_records = []
        self.nmea_path = ""
        self._build_ui()

    def _build_ui(self):
        root = QWidget()
        layout = QVBoxLayout(root)
        self.drop = DropArea()
        self.drop.on_files = self.add_paths
        layout.addWidget(self.drop)

        buttons = QHBoxLayout()
        for text, handler in [
            ("ファイル追加", self.choose_files),
            ("フォルダ追加", self.choose_folder),
            ("NMEA読込", self.choose_nmea),
            ("NMEA照合", self.match_nmea),
            ("クリア", self.clear),
            ("TXT出力", self.export_txt),
            ("CSV出力", self.export_csv),
            ("Excel出力", self.export_xlsx),
            ("JSON全タグ出力", self.export_json),
        ]:
            btn = QPushButton(text)
            btn.clicked.connect(handler)
            buttons.addWidget(btn)
        buttons.addWidget(QLabel("照合±秒"))
        self.threshold = QDoubleSpinBox()
        self.threshold.setDecimals(2)
        self.threshold.setRange(0.01, 60.0)
        self.threshold.setValue(0.5)
        self.threshold.setSingleStep(0.1)
        buttons.addWidget(self.threshold)
        buttons.addWidget(QLabel("NMEA UTC補正"))
        self.utc_offset = QSpinBox()
        self.utc_offset.setRange(-12, 14)
        self.utc_offset.setValue(9)
        buttons.addWidget(self.utc_offset)
        self.recursive_note = QCheckBox("フォルダはサブフォルダも検索")
        self.recursive_note.setChecked(True)
        self.recursive_note.setEnabled(False)
        buttons.addWidget(self.recursive_note)
        buttons.addStretch()
        layout.addLayout(buttons)

        self.progress = QProgressBar()
        self.progress.setVisible(False)
        layout.addWidget(self.progress)

        splitter = QSplitter(Qt.Vertical)
        self.table = QTableView()
        self.table.setModel(self.model)
        self.table.clicked.connect(self.show_raw_for_row)
        splitter.addWidget(self.table)
        self.detail = QTextEdit()
        self.detail.setReadOnly(True)
        splitter.addWidget(self.detail)
        splitter.setSizes([520, 230])
        layout.addWidget(splitter)
        self.status = QLabel("準備完了")
        layout.addWidget(self.status)
        self.setCentralWidget(root)

        help_action = QAction("補足", self)
        help_action.triggered.connect(self.show_help)
        self.menuBar().addAction(help_action)

    def show_help(self):
        QMessageBox.information(
            self, "補足",
            "NMEA時刻はUTCとして読み込み、標準で+9時間してカメラ撮影時刻と照合します。\n"
            "GGAは位置と精度系、PRDIDは第1値から第3値をそのまま保持します。\n"
            "PRDIDの意味は機器設定により異なる場合があるため、最初はPRDID_1/2/3として確認してください。"
        )

    def choose_files(self):
        files, _ = QFileDialog.getOpenFileNames(self, "PEFファイルを選択", "", "PEF files (*.pef *.PEF);;All files (*.*)")
        if files:
            self.add_paths(files)

    def choose_folder(self):
        folder = QFileDialog.getExistingDirectory(self, "PEFフォルダを選択")
        if folder:
            self.add_paths([folder])

    def add_paths(self, paths):
        files = collect_pef_paths(paths)
        if not files:
            QMessageBox.warning(self, "PEFなし", "PEFファイルが見つかりませんでした。")
            return
        self.analyze(files)

    def analyze(self, files):
        try:
            self.progress.setVisible(True)
            self.progress.setRange(0, 0)
            QApplication.processEvents()
            self.status.setText(f"PEF解析中: {len(files)}件")
            self.raw_print = run_exiftool(files, numeric=False)
            self.raw_numeric = run_exiftool(files, numeric=True)
            rows = []
            for i, pdata in enumerate(self.raw_print):
                ndata = self.raw_numeric[i] if i < len(self.raw_numeric) else {}
                rows.append(build_row(pdata, ndata))
            self.rows = rows
            if self.nmea_records:
                self.apply_nmea_matching()
            self.model.set_rows(self.rows)
            self.table.resizeColumnsToContents()
            self.status.setText(f"PEF解析完了: {len(self.rows)}件")
            self.detail.setPlainText("行をクリックすると全タグを表示します。")
        except Exception as e:
            QMessageBox.critical(self, "エラー", f"解析に失敗しました。\n\n{e}\n\n{traceback.format_exc()}")
            self.status.setText("エラー")
        finally:
            self.progress.setVisible(False)

    def choose_nmea(self):
        path, _ = QFileDialog.getOpenFileName(self, "NMEAファイルを選択", "", "NMEA/Text (*.nmea *.txt *.log);;All files (*.*)")
        if not path:
            return
        try:
            fallback = None
            for row in self.rows:
                dt = parse_exif_datetime(row.get("撮影時刻"))
                if dt:
                    fallback = dt.date()
                    break
            recs, bad = parse_nmea_file(path, self.utc_offset.value(), fallback)
            self.nmea_records = recs
            self.nmea_path = path
            self.status.setText(f"NMEA読込完了: {len(recs)} GGA記録 / checksum不一致 {bad} 行")
            if self.rows:
                self.apply_nmea_matching()
                self.model.set_rows(self.rows)
                self.table.resizeColumnsToContents()
            self.detail.setPlainText(f"NMEA: {path}\nGGA記録: {len(recs)}\nchecksum不一致: {bad}")
        except Exception as e:
            QMessageBox.critical(self, "NMEAエラー", f"NMEA読込に失敗しました。\n\n{e}\n\n{traceback.format_exc()}")

    def match_nmea(self):
        if not self.nmea_records:
            QMessageBox.warning(self, "NMEAなし", "先にNMEAファイルを読み込んでください。")
            return
        if not self.rows:
            QMessageBox.warning(self, "PEFなし", "先にPEFを読み込んでください。")
            return
        self.apply_nmea_matching()
        self.model.set_rows(self.rows)
        self.table.resizeColumnsToContents()
        ok = sum(1 for r in self.rows if r.get("NMEA照合") == "OK")
        self.status.setText(f"NMEA照合完了: OK {ok}/{len(self.rows)}件")

    def apply_nmea_matching(self):
        th = self.threshold.value()
        for row in self.rows:
            for c in NMEA_COLUMNS:
                row[c] = ""
            dt = parse_exif_datetime(row.get("撮影時刻"))
            rec, diff = nearest_nmea_record(self.nmea_records, dt, th)
            if rec is None:
                row["NMEA照合"] = "NG"
                row["時刻差秒"] = diff if diff != "" else ""
                continue
            row["NMEA照合"] = "OK"
            row["NMEA時刻"] = rec["dt_local"].strftime("%Y:%m:%d %H:%M:%S.%f")[:-3]
            row["時刻差秒"] = round(diff, 3)
            row["NMEA緯度"] = rec.get("lat", "")
            row["NMEA経度"] = rec.get("lon", "")
            row["NMEA標高"] = rec.get("altitude", "")
            row["GGA測位品質"] = rec.get("quality", "")
            row["GGA衛星数"] = rec.get("satellites", "")
            row["GGA_HDOP"] = rec.get("hdop", "")
            row["GGAジオイド高"] = rec.get("geoid_sep", "")
            row["GGA補正経過秒"] = rec.get("dgps_age", "")
            row["GGA基準局ID"] = rec.get("station_id", "")
            row["PRDID_1"] = rec.get("prdid_1", "")
            row["PRDID_2"] = rec.get("prdid_2", "")
            row["PRDID_3"] = rec.get("prdid_3", "")
            row["PRDID文"] = rec.get("prdid_sentence", "")
            d = haversine_m(row.get("緯度"), row.get("経度"), rec.get("lat"), rec.get("lon"))
            row["水平距離差m"] = round(d, 3) if d != "" else ""
            try:
                row["標高差m"] = round(float(row.get("標高")) - float(rec.get("altitude")), 3)
            except Exception:
                row["標高差m"] = ""

    def show_raw_for_row(self, index):
        r = index.row()
        lines = []
        if 0 <= r < len(self.rows):
            lines.append("[出力行]")
            lines.append(json.dumps(self.rows[r], ensure_ascii=False, indent=2))
        if 0 <= r < len(self.raw_print):
            lines.append("\n[EXIF/MakerNotes 全タグ]")
            lines.append(json.dumps(self.raw_print[r], ensure_ascii=False, indent=2))
        self.detail.setPlainText("\n".join(lines))

    def clear(self):
        self.raw_print = []
        self.raw_numeric = []
        self.rows = []
        self.nmea_records = []
        self.nmea_path = ""
        self.model.set_rows([])
        self.detail.clear()
        self.status.setText("クリアしました")

    def ensure_rows(self):
        if not self.rows:
            QMessageBox.warning(self, "データなし", "先にPEFを読み込んでください。")
            return False
        return True

    def default_name(self, ext):
        dt = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"pef_exif_nmea_{dt}.{ext}"

    def export_txt(self):
        if not self.ensure_rows():
            return
        path, _ = QFileDialog.getSaveFileName(self, "TXT保存", self.default_name("txt"), "Text (*.txt)")
        if not path:
            return
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            for row in self.rows:
                f.write("=" * 70 + "\n")
                for c in OUTPUT_COLUMNS:
                    f.write(f"{c}: {row.get(c, '')}\n")
        self.status.setText(f"TXT出力: {path}")

    def export_csv(self):
        if not self.ensure_rows():
            return
        path, _ = QFileDialog.getSaveFileName(self, "CSV保存", self.default_name("csv"), "CSV (*.csv)")
        if not path:
            return
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=OUTPUT_COLUMNS)
            writer.writeheader()
            writer.writerows(self.rows)
        self.status.setText(f"CSV出力: {path}")

    def export_xlsx(self):
        if not self.ensure_rows():
            return
        if Workbook is None:
            QMessageBox.critical(self, "openpyxlなし", "openpyxlがインストールされていません。")
            return
        path, _ = QFileDialog.getSaveFileName(self, "Excel保存", self.default_name("xlsx"), "Excel (*.xlsx)")
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "PEF_EXIF_NMEA"
        ws.append(OUTPUT_COLUMNS)
        for row in self.rows:
            ws.append([row.get(c, "") for c in OUTPUT_COLUMNS])
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx, col in enumerate(OUTPUT_COLUMNS, start=1):
            max_len = len(col)
            for row_idx in range(2, ws.max_row + 1):
                v = ws.cell(row=row_idx, column=col_idx).value
                max_len = max(max_len, len(str(v)) if v is not None else 0)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 45)
        wb.save(path)
        self.status.setText(f"Excel出力: {path}")

    def export_json(self):
        if not self.raw_print:
            QMessageBox.warning(self, "データなし", "先にPEFを読み込んでください。")
            return
        path, _ = QFileDialog.getSaveFileName(self, "JSON保存", self.default_name("json"), "JSON (*.json)")
        if not path:
            return
        payload = {"rows": self.rows, "exif_raw": self.raw_print, "nmea_file": self.nmea_path, "nmea_records": self.nmea_records}
        def conv(o):
            if isinstance(o, datetime):
                return o.isoformat(sep=" ")
            return str(o)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2, default=conv)
        self.status.setText(f"JSON出力: {path}")


def main():
    app = QApplication(sys.argv)
    w = MainWindow()
    w.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
